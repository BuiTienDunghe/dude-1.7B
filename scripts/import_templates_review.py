"""Nhập lại file export/templates_review.xlsx (đã điền) vào lớp ghi đè QUA API của công cụ đang chạy (POST /api/template, có token + Host),
không nối thẳng vào edits_log (STATE trong bộ nhớ công cụ sẽ lệch). Công cụ không chạy thì dừng, không ghi gì, không tự khởi động lại.

  OK                          -> "duyet" (chuỗi chưa duyệt thành đã duyệt). Dòng đã duyệt hoặc đã sửa: bỏ qua (không đổi).
  SUA + cột "tiếng Việt sửa"  -> "sua" (công cụ đặt trạng thái "đã sửa"; release_check coi "đã sửa" là đã xử lý, chỉ chặn "chưa duyệt").
                                 Chuỗi bằng giá trị hiện hành: chưa duyệt -> "duyet"; đã duyệt/đã sửa -> bỏ qua.
  trống                       -> giữ nguyên; liệt kê các chuỗi còn chưa duyệt.
Kiểm chuỗi sửa: không rỗng, NFC (tự chuẩn hóa nếu chưa, có báo), giữ nguyên placeholder {q}/{claim}... như chuỗi gốc. Ghi chú đưa vào ghi_chu.
Idempotent: chạy lại cùng file không ghi bản ghi mới. Cuối cùng chạy release_check và in kết quả.

Dùng (Python Windows, thư mục dự án):
  python scripts/import_templates_review.py --xlsx export/templates_review.xlsx            (ghi vào công cụ ở cổng 8765)
  python scripts/import_templates_review.py --xlsx export/templates_review.xlsx --dry_run  (chỉ in kế hoạch)
"""
import argparse
import json
import re
import subprocess
import sys
import unicodedata
import urllib.request
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools/review_app"))
import edits as E  # noqa: E402
from inventory import inventory  # noqa: E402

for _s in (sys.stdout, sys.stderr):
    try:
        _s.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass
ap = argparse.ArgumentParser()
ap.add_argument("--xlsx", default=str(ROOT / "export/templates_review.xlsx"))
ap.add_argument("--port", type=int, default=8765)
ap.add_argument("--edits_dir", default=str(ROOT / "data/human_edits"), help="lớp ghi đè mà công cụ đang dùng (để đọc trạng thái hiện hành)")
ap.add_argument("--config_dir", default=str(ROOT / "config"))
ap.add_argument("--dry_run", action="store_true")
args = ap.parse_args()
HOST = f"127.0.0.1:{args.port}"
PH = re.compile(r"\{[a-z_]+\}")

# ---------------------------------------------------------------- đọc file
ws = load_workbook(args.xlsx).active
hr = next(r for r in range(1, 30) if ws.cell(row=r, column=1).value == "khóa")
col = {ws.cell(row=hr, column=j).value: j for j in range(1, ws.max_column + 1) if ws.cell(row=hr, column=j).value}
for need in ("khóa", "tiếng Việt sửa", "quyết định", "ghi chú"):
    assert need in col, f"thiếu cột {need}"
get = lambda r, h: ws.cell(row=r, column=col[h]).value
inv = {x["key"]: x for x in inventory(args.config_dir)}
st = E.derive_state(E.read_log(Path(args.edits_dir) / "edits_log.jsonl"), list(inv.values()))

plan, skipped, blank, errors, nfc_fixed, seen = [], [], [], [], [], set()
for r in range(hr + 1, ws.max_row + 1):
    key = get(r, "khóa")
    if not key:
        continue
    key = str(key).strip()
    if key not in inv:
        errors.append(f"dòng {r}: khóa lạ {key}"); continue
    if key in seen:
        errors.append(f"dòng {r}: khóa lặp {key}"); continue
    seen.add(key)
    m = st["mau"][key]
    dec = str(get(r, "quyết định") or "").strip().upper()
    note = str(get(r, "ghi chú") or "").strip()
    new = get(r, "tiếng Việt sửa")
    new = "" if new is None else str(new).strip()
    if not dec:
        if m["trang_thai"] == "chưa duyệt":
            blank.append(key)
        continue
    if dec == "OK":
        if m["trang_thai"] == "chưa duyệt":
            plan.append({"key": key, "action": "duyet", "value": "", "note": note})
        else:
            skipped.append((key, f"OK, đã {m['trang_thai'].split()[-1]}"))
    elif dec == "SUA":
        if not new:
            errors.append(f"dòng {r}: {key} chọn SUA nhưng cột 'tiếng Việt sửa' trống"); continue
        if not unicodedata.is_normalized("NFC", new):
            new = unicodedata.normalize("NFC", new); nfc_fixed.append(key)
        if sorted(PH.findall(new)) != sorted(PH.findall(m["goc"])):
            errors.append(f"dòng {r}: {key} chuỗi sửa không giữ placeholder {sorted(PH.findall(m['goc']))}"); continue
        if new == m["gia_tri"]:
            if m["trang_thai"] == "chưa duyệt":
                plan.append({"key": key, "action": "duyet", "value": "", "note": note})
            else:
                skipped.append((key, "SUA, chuỗi bằng giá trị hiện hành"))
        else:
            plan.append({"key": key, "action": "sua", "value": new, "note": note})
    else:
        errors.append(f"dòng {r}: {key} quyết định không hợp lệ '{dec}' (chỉ OK / SUA / trống)")
missing_rows = sorted(set(inv) - seen)

print(f"File: {args.xlsx}")
print(f"Kế hoạch: duyệt (OK) {sum(p['action'] == 'duyet' for p in plan)}, sửa {sum(p['action'] == 'sua' for p in plan)}; "
      f"bỏ qua vì không đổi {len(skipped)}; còn trống {len(blank)}; lỗi {len(errors)}; tự chuẩn hóa NFC {len(nfc_fixed)}.")
for x in errors:
    print("  LỖI:", x)
if missing_rows:
    print(f"  Chú ý: {len(missing_rows)} khóa không có trong file: {', '.join(missing_rows)}")
if errors:
    print("Có lỗi: KHÔNG ghi gì. Sửa file rồi chạy lại.")
    sys.exit(2)
if args.dry_run:
    for p in plan:
        print(f"  {p['action']:5} {p['key']}" + (f" -> {p['value']}" if p["action"] == "sua" else ""))
    sys.exit(0)

# ---------------------------------------------------------------- ghi qua API công cụ
try:
    html = urllib.request.urlopen(urllib.request.Request(f"http://{HOST}/", headers={"Host": HOST}), timeout=10).read().decode("utf-8")
    token = json.loads(re.search(r'var TOKEN=("[^"]+")', html).group(1))
except Exception as ex:
    print(f"Công cụ không chạy ở {HOST} ({type(ex).__name__}). KHÔNG ghi gì; báo phiên điều phối (đừng tự khởi động lại).")
    sys.exit(3)
# chốt chặn: công cụ ở cổng này phải đang dùng đúng lớp ghi đè --edits_dir (so số bản ghi công cụ báo với số dòng edits_log)
m_ = re.search(r"Bản ghi trong log hiệu đính: (\d+)", html)
n_log = len(E.read_log(Path(args.edits_dir) / "edits_log.jsonl"))
if not m_ or int(m_.group(1)) != n_log:
    print(f"Công cụ ở {HOST} báo {m_.group(1) if m_ else '?'} bản ghi, còn {args.edits_dir} có {n_log}: công cụ này KHÔNG dùng lớp ghi đè đã chỉ định. "
          "KHÔNG ghi gì; báo phiên điều phối.")
    sys.exit(4)
done = 0
for p in plan:
    body = json.dumps({"key": p["key"], "action": p["action"], "value": p["value"], "note": ("[excel] " + p["note"]).strip()}).encode("utf-8")
    req = urllib.request.Request(f"http://{HOST}/api/template", data=body, method="POST",
                                 headers={"Host": HOST, "Content-Type": "application/json", "X-Review-Token": token})
    r = json.loads(urllib.request.urlopen(req, timeout=30).read().decode("utf-8"))
    assert r.get("ok"), (p["key"], r)
    done += 1
n_after = len(E.read_log(Path(args.edits_dir) / "edits_log.jsonl"))
assert n_after == n_log + done, f"số bản ghi trong {args.edits_dir} tăng {n_after - n_log}, mong đợi {done}"
st2 = E.derive_state(E.read_log(Path(args.edits_dir) / "edits_log.jsonl"), list(inv.values()))
print(f"Đã ghi {done} bản ghi qua công cụ. OK {sum(p['action'] == 'duyet' for p in plan)}, sửa {sum(p['action'] == 'sua' for p in plan)}, "
      f"còn trống {len(blank)}, bỏ qua vì không đổi {len(skipped)}.")
if blank:
    print("Còn chưa duyệt (dòng trống):", ", ".join(blank))
print("Trạng thái chuỗi mẫu sau khi nhập:", st2["dem_mau"])
rc = subprocess.run([sys.executable, str(ROOT / "scripts/human_edits/release_check.py"), "--edits_dir", args.edits_dir, "--config_dir", args.config_dir],
                    capture_output=True, text=True, encoding="utf-8", env={**__import__("os").environ, "PYTHONIOENCODING": "utf-8"})
print(f"release_check (mã {rc.returncode}):")
print("\n".join([l for l in rc.stdout.splitlines() if not l.startswith("  ")][:6]))
