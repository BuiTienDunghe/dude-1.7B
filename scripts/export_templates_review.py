"""Xuất 78 chuỗi mẫu tiếng Việt (tools/review_app/inventory.py + trạng thái hiện hành của lớp ghi đè) ra export/templates_review.xlsx để người làm
dự án duyệt trong Excel, rồi nhập lại bằng scripts/import_templates_review.py. Một sheet; chỉ 3 cột "tiếng Việt sửa", "quyết định", "ghi chú" được sửa
(bảo vệ sheet, không mật khẩu); "quyết định" là danh sách thả xuống OK / SUA. Dòng đã duyệt điền sẵn OK, dòng đã sửa điền sẵn SUA + chuỗi hiện hành.
Python Windows (openpyxl)."""
import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools/review_app"))
import edits as E  # noqa: E402
from inventory import inventory  # noqa: E402

ap = argparse.ArgumentParser()
ap.add_argument("--out", default=str(ROOT / "export/templates_review.xlsx"))
ap.add_argument("--edits_dir", default=str(ROOT / "data/human_edits"))
ap.add_argument("--config_dir", default=str(ROOT / "config"))
args = ap.parse_args()

GROUP = {"Câu hỏi mẫu": "câu hỏi mẫu", "Mô tả ý định MASSIVE": "ý định MASSIVE", "Tên chủ đề SIB-200": "chủ đề SIB",
         "Lựa chọn kiểm chứng": "lựa chọn kiểm chứng", "Có/Không": "Có-Không"}
ORDER = ["câu hỏi mẫu", "ý định MASSIVE", "chủ đề SIB", "lựa chọn kiểm chứng", "Có-Không"]
HEAD = ["khóa", "nhóm", "tiếng Anh", "tiếng Việt hiện tại", "tiếng Việt sửa", "quyết định", "ghi chú", "trạng thái hiện tại", "gốc (trước khi sửa)"]
EDITABLE = {"tiếng Việt sửa", "quyết định", "ghi chú"}
WIDTH = {"khóa": 44, "nhóm": 18, "tiếng Anh": 42, "tiếng Việt hiện tại": 46, "tiếng Việt sửa": 46, "quyết định": 12, "ghi chú": 36,
         "trạng thái hiện tại": 16, "gốc (trước khi sửa)": 40}

inv = inventory(args.config_dir)
unknown = sorted({x["group"] for x in inv} - set(GROUP))
assert not unknown, f"nhóm inventory chưa ánh xạ: {unknown}"
st = E.derive_state(E.read_log(Path(args.edits_dir) / "edits_log.jsonl"), inv)


def sort_key(x):
    g = GROUP[x["group"]]
    return (ORDER.index(g), x["key"].rsplit(".", 1)[-1] if g == "ý định MASSIVE" else x["key"])


items = sorted(inv, key=sort_key)
FONT, BOLD = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
FILL = {"chưa duyệt": None, "đã duyệt": PatternFill("solid", fgColor="E8F5E9"), "đã sửa": PatternFill("solid", fgColor="FFF8E1")}
EDIT_FILL = PatternFill("solid", fgColor="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="DDE3EA")
thin = Side(style="thin", color="C8CDD3")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws = wb.active
ws.title = "chuoi_mau"
legend = [
    "Duyệt 78 chuỗi mẫu tiếng Việt của benchmark (câu hỏi mẫu, mô tả ý định MASSIVE, chủ đề SIB, lựa chọn kiểm chứng, Có/Không).",
    "CHỈ ĐIỀN 3 cột có nền trắng viền đậm: \"tiếng Việt sửa\", \"quyết định\" (chọn OK hoặc SUA trong danh sách), \"ghi chú\". Các cột khác khóa chỉ đọc.",
    "OK = chuỗi hiện tại đúng, duyệt. SUA = dùng chuỗi ở cột \"tiếng Việt sửa\" (bắt buộc điền). Để trống = chưa quyết, giữ trạng thái chưa duyệt.",
    "Nền xanh = đã duyệt (điền sẵn OK); nền vàng = đã sửa (điền sẵn SUA + chuỗi hiện hành, cột cuối là chuỗi gốc). Giữ nguyên {q}, {claim} nếu chuỗi gốc có.",
]
for i, t in enumerate(legend, 1):
    c = ws.cell(row=i, column=1, value=t)
    c.font = BOLD if i == 1 else FONT
HR = len(legend) + 2
for j, h in enumerate(HEAD, 1):
    c = ws.cell(row=HR, column=j, value=h)
    c.font, c.fill, c.border = BOLD, HEAD_FILL, BORDER
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[c.column_letter].width = WIDTH[h]
dv = DataValidation(type="list", formula1='"OK,SUA"', allow_blank=True, showErrorMessage=True, errorTitle="Quyết định",
                    error="Chỉ chọn OK hoặc SUA (hoặc để trống).")
ws.add_data_validation(dv)
col = {h: j for j, h in enumerate(HEAD, 1)}
for i, x in enumerate(items, HR + 1):
    m = st["mau"][x["key"]]
    status = m["trang_thai"]
    vals = {"khóa": x["key"], "nhóm": GROUP[x["group"]], "tiếng Anh": x["en"], "tiếng Việt hiện tại": m["gia_tri"],
            "tiếng Việt sửa": m["gia_tri"] if status == "đã sửa" else None, "quyết định": {"đã duyệt": "OK", "đã sửa": "SUA"}.get(status),
            "ghi chú": None, "trạng thái hiện tại": status, "gốc (trước khi sửa)": m["goc"] if status == "đã sửa" else None}
    for h, v in vals.items():
        c = ws.cell(row=i, column=col[h], value=v)
        c.font, c.border = FONT, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if h in EDITABLE:
            c.protection = Protection(locked=False)
            c.fill = EDIT_FILL
            c.border = Border(left=Side(style="medium", color="5B6B7C"), right=Side(style="medium", color="5B6B7C"),
                              top=thin, bottom=thin)
        elif FILL[status] is not None:
            c.fill = FILL[status]
    dv.add(ws.cell(row=i, column=col["quyết định"]))
last = HR + len(items)
ws.freeze_panes = ws.cell(row=HR + 1, column=1)
ws.auto_filter.ref = f"A{HR}:{ws.cell(row=last, column=len(HEAD)).column_letter}{last}"
ws.protection.sheet = True              # không mật khẩu
ws.protection.autoFilter = False        # vẫn lọc được khi sheet bị bảo vệ
ws.protection.sort = False
ws.protection.formatColumns = False
ws.protection.formatRows = False
Path(args.out).parent.mkdir(parents=True, exist_ok=True)
wb.save(args.out)
from collections import Counter  # noqa: E402
print({"file": args.out, "dong": len(items), "trang_thai": dict(Counter(st["mau"][x["key"]]["trang_thai"] for x in items)),
       "nhom": dict(Counter(GROUP[x["group"]] for x in items)), "hang_tieu_de": HR})
